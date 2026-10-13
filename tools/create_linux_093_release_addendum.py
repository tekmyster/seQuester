from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor

from create_linux_092_release_addendum import (
    BLUE,
    NAVY,
    RED,
    add_bullet,
    add_note,
    add_table,
    configure_section,
    configure_styles,
)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--logo", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    document = Document()
    configure_styles(document)
    configure_section(document.sections[0])
    header = document.sections[0].header.paragraphs[0]
    header.text = "Quiver seQuester  |  Linux 0.9.3 release status addendum"
    footer = document.sections[0].footer.paragraphs[0]
    for run in footer.runs:
        if "RF DOC 001-A" in run.text:
            run.text = run.text.replace("RF DOC 001-A", "RF DOC 001-B")

    logo_p = document.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo = logo_p.add_run().add_picture(args.logo, width=Inches(2.0))
    doc_pr = logo._inline.xpath("./wp:docPr")[0]
    doc_pr.set("descr", "Quiver seQuester shield and document logo")
    doc_pr.set("title", "Quiver seQuester logo")
    p = document.add_paragraph("Quiver product line", style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = document.add_paragraph("Linux x64 0.9.3 Release Status Addendum", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = document.add_paragraph("Current addendum to seQuester Product Manual, Revision 1.2", style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Published 13 September 2026  •  Unsigned development release  •  Fleet rollout held")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(RED)
    document.add_paragraph()
    add_note(
        document,
        "Supersession",
        "This addendum supersedes the Linux activation guidance in RF DOC 001-A (the 0.9.2 addendum). "
        "Revision 1.2 and prior addenda remain retained as dated records; use this document with the current packaged audit guide and pilot procedure.",
    )
    add_note(
        document,
        "Acceptance boundary",
        "The 0.9.3 archive and isolated launcher fixture are verified. Full live canary attribution, the same EventId in management and Wazuh, setgid packet permissions, reboot/recovery and UI acceptance remain pending.",
        danger=True,
    )
    document.add_paragraph("Document control", style="Heading 1")
    add_table(
        document,
        ["Field", "Value"],
        [
            ["Document identifier", "RF DOC 001-B"],
            ["Supersedes", "RF DOC 001-A activation guidance only"],
            ["Product", "Quiver seQuester"],
            ["Scope", "Linux x64 endpoint audit-attribution package 0.9.3"],
            ["Status", "Development only; live pilot pending; fleet rollout held"],
        ],
        [1.55, 5.8],
    )

    document.add_page_break()
    document.add_paragraph("1. Published package", style="Heading 1")
    add_table(
        document,
        ["Property", "Verified value"],
        [
            ["Release", "ransomfree-client-0.9.3-dev.20260913T170332Z-dec6cf70-linux-x64"],
            ["Archive", "ransomfree-client-0.9.3-dev.20260913T170332Z-dec6cf70-linux-x64.tar.gz"],
            ["SHA-256", "5d27ed74d40b0c251b04ff8a2b3c2bf939b55839a2ebd1fd62b9b13ddbfa1b30"],
            ["Size and inventory", "85,036,258 bytes; 476 files verified"],
            ["Endpoint / relay", "Endpoint 0.9.3; unchanged Native AOT relay component 0.9.1, protocol 1"],
            ["Approval", "Unsigned development candidate; not production or fleet approved"],
        ],
        [1.55, 5.8],
    )
    document.add_paragraph("2. Corrected zero-argument launcher", style="Heading 1")
    p = document.add_paragraph()
    p.add_run("auditd → ").bold = True
    p.add_run("root-owned launch-relay.sh → setpriv → same-PID Native AOT relay → private AF_UNIX socket → root collector → protected inbox")
    add_note(
        document,
        "Required auditd plugin",
        "Set path to /opt/ransomfree-client/current/audit/launch-relay.sh and omit the args directive entirely. "
        "The validated launcher contains all four setpriv capability/no-new-privileges flags and the final relay path.",
    )
    add_note(
        document,
        "Do not reuse 0.9.2 configuration",
        "Auditd 3.0.7 permits at most two plugin arguments. The 0.9.2 five-token setpriv line failed before the relay connected or a canary was read. Do not configure setpriv directly, use empty args, or retain the five-token line.",
        danger=True,
    )
    document.add_paragraph("3. Launcher and credential contract", style="Heading 1")
    add_table(
        document,
        ["Control", "Required state"],
        [
            ["Launcher", "Root-owned mode 0755; no arguments during auditd execution; --check is a manual no-start preflight only."],
            ["Release path", "Root-owned canonical versioned release and non-writable executables/ancestors; current is a root-owned selector symlink."],
            ["Collector", "User=root and Group=root; no capabilities. Remove conflicting old Group=rfclient overrides through reviewed maintenance."],
            ["Relay", "Root UID/GID, CapPrm/CapEff/CapBnd=0, NoNewPrivs=1; final executable identity and auditd parent retained in the same PID."],
            ["Inbox", "root:rfclient mode 2750; setgid inheritance preserves endpoint-readable packet group ownership."],
        ],
        [1.45, 5.9],
    )

    document.add_page_break()
    document.add_paragraph("4. Isolated native fixture evidence", style="Heading 1")
    add_table(
        document,
        ["Assertion", "Verified result"],
        [
            ["Execution environment", "MDWE=1 applied and verified; zero plugin arguments; NULL environment; root UID/GID for launcher/collector path."],
            ["Privilege state", "Collector and relay capability sets zero; NoNewPrivileges=1; no host mount, network or privileged/capability-added container."],
            ["Identity continuity", "Launcher exec preserved the PID and final relay executable identity; auditd-parent contract retained."],
            ["Transport", "88 input bytes; zero failures."],
            ["Negative cases", "Rejected arguments, non-root caller, writable launcher, writable relay and writable release ancestors."],
        ],
        [1.75, 5.6],
    )
    add_note(
        document,
        "Evidence limit",
        "The packaged JSON proves the isolated launcher/transport contract only. It did not activate the live host's auditd, read a real canary or contact management/Wazuh.",
        danger=True,
    )
    document.add_paragraph("5. Shipped rule-loading boundary", style="Heading 1")
    add_bullet(document, "--audit-rules remains a reviewable shell-command draft, not a persistent rules file.")
    add_bullet(document, "--audit-rules-file emits direct auditctl -R syntax only for restricted ASCII-safe paths; richer names fail closed and need separate reviewed handling.")
    add_bullet(document, "sequester-audit-rules.service reads an administrator-reviewed root:root mode-0600 file outside /etc/audit/rules.d; automatic reconciliation, duplicate-free reload and reboot persistence are not promised.")
    document.add_paragraph("6. Current acceptance state", style="Heading 1")
    add_table(
        document,
        ["Gate", "Status", "Required evidence"],
        [
            ["Package integrity", "Verified", "Exact archive hash, size and 476-file inventory."],
            ["Isolated launcher fixture", "Passed", "Positive and negative launcher/privilege/transport assertions recorded in packaged JSON."],
            ["Live controlled canary attribution", "Pending", "Fresh full-attribution event with matching PID/PPID/UID/eUID/executable and unchanged canary content."],
            ["Setgid packet permissions", "Pending", "Collector-created packets retain the expected rfclient-readable group and protected ownership."],
            ["Management and Wazuh", "Pending", "The exact attributed EventId is retained by endpoint, management and Wazuh."],
            ["Reboot/recovery and UI", "Pending", "Service/rule persistence, bounded failure/recovery, rollback and operator status presentation."],
            ["Fleet rollout", "Held", "All live pilot gates approved by the rollout owner."],
        ],
        [2.05, 1.15, 4.15],
    )
    add_note(
        document,
        "Rollback",
        "If a live gate fails, stop expansion, preserve diagnostics/inbox/queue, remove only scoped rules and return to healthy detection-only 0.9.0. Never reactivate the old JIT plugin or weaken auditd MDWE/ProtectHome.",
        danger=True,
    )
    document.add_paragraph("7. Current references", style="Heading 1")
    for item in (
        "ATTRIBUTION_RELEASE.md",
        "packaging/linux-client/AUDIT.md",
        "packaging/linux-client/AUDIT_PILOT.md",
        "CLIENT_FEATURE_MATRIX.md",
        "DOCUMENTATION_CHANGE_LEDGER.md",
        "QA.md",
    ):
        add_bullet(document, item)

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Created {output}")


if __name__ == "__main__":
    main()
