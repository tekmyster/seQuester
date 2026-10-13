"""Independent parser checks for the embedded templates; no authoring."""
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

assets = Path(__file__).resolve().parents[1] / "Assets"
doc = Document(assets / "canary.docx")
assert doc.paragraphs[0].text == "Account recovery register"
assert doc.paragraphs[0].style.name == "Title"
assert any("Synthetic security canary" in p.text for p in doc.paragraphs)
workbook = load_workbook(assets / "canary.xlsx", read_only=True, data_only=True)
assert workbook.sheetnames == ["Recovery register"]
assert workbook.active["D5"].value == 2
assert workbook.active["D6"].value == 1
workbook.close()
pdf = PdfReader(assets / "canary.pdf", strict=True)
assert len(pdf.pages) == 1
assert "Confidential archive inventory" in pdf.pages[0].extract_text()
for name in ["canary.docx", "canary.xlsx"]:
    with zipfile.ZipFile(assets / name) as package:
        assert package.testzip() is None
        for item in package.namelist():
            assert not item.lower().endswith("vbaproject.bin")
            if item.endswith(".xml") or item.endswith(".rels"):
                root = ET.fromstring(package.read(item))
                assert all(element.attrib.get("TargetMode") != "External" for element in root.iter())
print("PASS: DOCX, XLSX and PDF independently parsed; Office packages have no macros or external relationships.")
