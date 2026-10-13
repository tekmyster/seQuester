from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "173B63"
BLUE = "0B6EB7"
PALE = "EAF2F8"
RED = "D83A32"
GRAY = "5A6673"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    tr_pr = header._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for index, value in enumerate(headers):
        shade(header.cells[index], NAVY)
        set_cell_text(header.cells[index], value, True, "FFFFFF")
        if widths:
            header.cells[index].width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, value in enumerate(values):
            if row_index % 2:
                shade(row.cells[index], PALE)
            set_cell_text(row.cells[index], value)
            if widths:
                row.cells[index].width = Inches(widths[index])
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_field(run, instruction: str, cached: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = cached
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, value, end):
        run._r.append(node)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("1E2832")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in (
        ("Title", 30, NAVY),
        ("Subtitle", 16, BLUE),
        ("Heading 1", 19, NAVY),
        ("Heading 2", 13, NAVY),
    ):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def configure_section(section) -> None:
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)
    header = section.header
    p = header.paragraphs[0]
    p.text = "Quiver seQuester  |  Linux release status addendum"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.runs[0].font.name = "Aptos"
    p.runs[0].font.size = Pt(8)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("RF DOC 001-A  •  13 September 2026  |  Page ")
    add_field(p.add_run(), " PAGE ", "1")
    p.add_run(" of ")
    add_field(p.add_run(), " NUMPAGES ", "3")
    for run in p.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)


def add_bullet(document: Document, text: str) -> None:
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def add_note(document: Document, label: str, text: str, danger: bool = False) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    cell = table.cell(0, 0)
    shade(cell, "FDEDEC" if danger else PALE)
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(f"{label}: ")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(RED if danger else NAVY)
    p.add_run(text)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--logo", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    document = Document()
    configure_styles(document)
    section = document.sections[0]
    configure_section(section)

    logo_paragraph = document.add_paragraph()
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo = logo_paragraph.add_run().add_picture(args.logo, width=Inches(2.05))
    doc_pr = logo._inline.xpath("./wp:docPr")[0]
    doc_pr.set("descr", "Quiver seQuester shield and document logo")
    doc_pr.set("title", "Quiver seQuester logo")
    document.add_paragraph("Quiver product line", style="Subtitle").alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = document.add_paragraph("Linux x64 0.9.2 Release Status Addendum", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Addendum to seQuester Product Manual, Revision 1.2", style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Published 13 September 2026  •  Development release  •  Fleet rollout held")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(RED)
    document.add_paragraph()
    add_note(
        document,
        "Purpose",
        "This dated addendum supersedes Revision 1.2 only for Linux x64 0.9.2 package and audit-attribution status. "
        "The complete manual remains authoritative for all other product areas.",
    )
    add_note(
        document,
        "Release boundary",
        "The package is unsigned development evidence. The isolated transport regression passed, but live auditd "
        "canary attribution, exact-event management/Wazuh delivery, persistence, recovery and UI acceptance remain pending.",
        danger=True,
    )
    document.add_paragraph("Document control", style="Heading 1")
    add_table(
        document,
        ["Field", "Value"],
        [
            ["Document identifier", "RF DOC 001-A"],
            ["Relationship", "Release-status addendum to RF DOC 001, Revision 1.2"],
            ["Product", "Quiver seQuester"],
            ["Scope", "Linux x64 endpoint audit-attribution package 0.9.2"],
            ["Status", "Development only; pilot pending; fleet rollout held"],
        ],
        [1.65, 5.7],
    )

    document.add_page_break()
    document.add_paragraph("1. Published package", style="Heading 1")
    add_table(
        document,
        ["Property", "Verified value"],
        [
            ["Release", "ransomfree-client-0.9.2-dev.20260913T164440Z-29942e83-linux-x64"],
            ["Archive", "ransomfree-client-0.9.2-dev.20260913T164440Z-29942e83-linux-x64.tar.gz"],
            ["SHA-256", "d10d1325f75ed21e365da0523dac8df43b2fe05fafec6e2e5752fa5c53bb5e7c"],
            ["Size", "85,033,466 bytes"],
            ["Inventory", "474 files verified"],
            ["Platform", "Linux x64/glibc endpoint; Linux ARM64 equivalent remains pending"],
            ["Approval", "Not production or fleet approved"],
        ],
        [1.35, 6.0],
    )

    document.add_paragraph("2. Corrected auditd launch architecture", style="Heading 1")
    p = document.add_paragraph()
    p.add_run("auditd → ").bold = True
    p.add_run("/usr/bin/setpriv → Native AOT relay → private AF_UNIX socket → managed collector → protected inbox → durable management queue")
    add_note(
        document,
        "Required plugin configuration",
        "Set path to /usr/bin/setpriv and provide a nonempty args line containing "
        "--bounding-set=-all --inh-caps=-all --ambient-caps=-all --no-new-privs followed by the absolute relay path. "
        "The relay itself accepts no arguments.",
    )
    add_note(
        document,
        "Important correction",
        "“Omit args entirely” applies only to the superseded direct no-argument relay configuration. It is wrong for "
        "the 0.9.2 setpriv plugin. An empty args directive remains invalid.",
        danger=True,
    )
    document.add_paragraph("3. Root cause and bounded proof", style="Heading 1")
    add_table(
        document,
        ["Fixture condition", "Observed result", "Interpretation"],
        [
            ["Capability-free collector; relay inherited a larger capability set", "PeerExecutableRead, errno 13 (EACCES)", "Collector could not complete the required /proc/PID/exe identity check."],
            ["Capability-free collector; setpriv-launched capability-free relay", "Transport succeeded; 88 input bytes; zero failures", "Capability mismatch correction confirmed without relaxing UID, executable or path identity checks."],
        ],
        [2.15, 2.2, 3.05],
    )
    add_bullet(document, "No CAP_SYS_PTRACE, host mount, network, container privilege or capability addition was used.")
    add_bullet(document, "The relay component intentionally reports version 0.9.1 and protocol 1 inside endpoint package 0.9.2.")
    add_bullet(document, "This fixture proves the transport correction only; it is not live auditd, canary, management or Wazuh acceptance.")

    document.add_page_break()
    document.add_paragraph("4. Shipped rule loading and diagnostics", style="Heading 1")
    add_table(
        document,
        ["Capability", "0.9.2 behavior", "Operator boundary"],
        [
            ["--audit-rules", "Produces a shell command draft for review.", "Never paste its shell quoting into an auditctl rules file."],
            ["--audit-rules-file", "Produces direct auditctl -R syntax without writing or applying it.", "Accepts only ASCII letters, digits, slash, underscore, dash and dot; fails closed for richer paths."],
            ["sequester-audit-rules.service", "Loads the reviewed rules file outside /etc/audit/rules.d, avoiding auditd ProtectHome visibility limits.", "Rules file must be administrator-reviewed, root:root and mode 0600. No automatic reconciliation or duplicate-free reload guarantee."],
            ["Transport diagnostics", "Records bounded failure count and fixed stage/category/errno/PID status.", "Does not record payloads, canary paths, command lines or exception text."],
        ],
        [1.65, 2.75, 3.0],
    )
    document.add_paragraph("5. Acceptance state", style="Heading 1")
    add_table(
        document,
        ["Gate", "Status", "Required evidence"],
        [
            ["Package integrity", "Verified", "Exact archive hash, byte count and 474-file inventory."],
            ["Isolated transport regression", "Passed", "Capability mismatch reproduced; setpriv correction passed with identity checks unchanged."],
            ["Live auditd and controlled canary attribution", "Pending", "A matching full-attribution event with expected PID/PPID/UID/eUID/executable and unchanged canary content."],
            ["Management and Wazuh", "Pending", "The exact attributed EventId retained by the endpoint, management server and Wazuh."],
            ["Persistence and recovery", "Pending", "Reboot, auditd restart, rule presence, duplicate prevention, bounded failure visibility and rollback."],
            ["UI/runtime acceptance", "Pending", "Current health/degraded state presentation and operator actions on the packaged client."],
            ["Fleet rollout", "Held", "All pilot gates approved by the rollout owner."],
        ],
        [2.15, 1.2, 4.05],
    )
    add_note(
        document,
        "Safe fallback",
        "On failure, stop expansion, preserve diagnostics/inbox/queue and return to the known detection-only 0.9.0 state. "
        "Do not reactivate the 0.9.0 JIT audit plugin or weaken auditd MemoryDenyWriteExecute/ProtectHome.",
        danger=True,
    )
    document.add_paragraph("6. Authoritative implementation references", style="Heading 1")
    for reference in (
        "packaging/linux-client/AUDIT.md",
        "packaging/linux-client/AUDIT_PILOT.md",
        "README.md",
        "CLIENT_FEATURE_MATRIX.md",
        "DOCUMENTATION_CHANGE_LEDGER.md",
        "QA.md",
    ):
        add_bullet(document, reference)

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Created {output}")


if __name__ == "__main__":
    main()
