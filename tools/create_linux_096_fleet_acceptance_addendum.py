from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from create_linux_092_release_addendum import GRAY, RED, add_field
from create_linux_095_release_addendum import set_note, set_row


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    document = Document(args.template)
    document.sections[0].header.paragraphs[0].text = "Quiver seQuester  |  Linux 0.9.6 fleet deployment acceptance"
    footer = document.sections[0].footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("RF DOC 001-H  •  13 September 2026  |  Page ")
    add_field(footer.add_run(), " PAGE ", "1")
    footer.add_run(" of ")
    add_field(footer.add_run(), " NUMPAGES ", "4")
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)

    replacements = {
        "Linux x64 0.9.6 Compatibility Acceptance Addendum": "Linux x64 0.9.6 Fleet Deployment Acceptance Addendum",
        "2. Accepted compatibility results": "2. Accepted fleet deployment results",
        "4. Verified native evidence": "4. Verified fleet evidence",
        "5. Remaining acceptance gates": "5. Acceptance status and remaining gates",
        "6. Controlled fleet expansion": "6. Operational boundary",
    }
    for paragraph in document.paragraphs:
        if paragraph.text in replacements:
            paragraph.text = replacements[paragraph.text]
        if "Two compatibility pilots accepted" in paragraph.text:
            paragraph.text = "Published 13 September 2026  •  Supported Linux fleet 26 of 26 accepted  •  Production certification pending"
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(RED)

    tables = document.tables
    set_note(
        tables[0].cell(0, 0),
        "Supersession",
        "This addendum supersedes RF DOC 001-G only for current supported Linux fleet deployment status. "
        "RF DOC 001-G remains the immutable two-pilot acceptance record, and all earlier addenda retain their historical scopes.",
    )
    set_note(
        tables[1].cell(0, 0),
        "Acceptance boundary",
        "The rollout owner reports deployment and final independent read-only validation passing on all 26 supported enrolled Linux endpoints, with zero changes, failures or unreachable targets within that denominator. "
        "This is supported-fleet deployment acceptance, not production certification or acceptance of excluded hosts and remaining platform gates.",
        True,
    )
    set_row(tables[2], 1, ["Document identifier", "RF DOC 001-H"])
    set_row(tables[2], 2, ["Supersedes", "RF DOC 001-G current fleet-status guidance only"])
    set_row(tables[2], 4, ["Scope", "Linux x64 0.9.6 supported enrolled fleet deployment acceptance"])
    set_row(tables[2], 5, ["Status", "Unsigned development candidate; supported Linux fleet 26/26 accepted; production pending"])

    package_rows = [
        ["Release", "ransomfree-client-0.9.6-dev.20260913T184639Z-08804ca0-linux-x64"],
        ["Archive", "ransomfree-client-0.9.6-dev.20260913T184639Z-08804ca0-linux-x64.tar.gz"],
        ["SHA-256", "501a4ec589b3af6f90c2eca6b4a7762b3f0c45c80961b21872310c567b77f02e"],
        ["Linux deployment", "26 of 26 supported enrolled endpoints accepted"],
        ["Canary preservation", "78 of 78 managed canary hashes preserved"],
        ["Management inventory", "28 acknowledged endpoints: 26 Linux 0.9.6 and 2 Windows 0.6.3"],
    ]
    for index, values in enumerate(package_rows, 1):
        set_row(tables[3], index, values)

    results = [
        "Every validated target reported 0.9.6-linux-dev with auditd, collector, reconciler and endpoint enabled and active.",
        "Every target retained six unique explicit-all rules, successful status bound to the current auditd PID, advancing health and an active zero-argument Native AOT plugin.",
        "Enrollment was present, automatic process-stop remained disabled, all 78 canary hashes were preserved and controlled reads produced Full attribution on all 26 hosts.",
        "Final independent read-only validation reported 26 of 26 passing with zero changes, failures or unreachable targets inside the supported enrolled set.",
    ]
    bullets = [p for p in document.paragraphs if p.style and p.style.name == "List Bullet"]
    for index, value in enumerate(results):
        bullets[index].text = value
    set_note(
        tables[4].cell(0, 0),
        "Final management evidence",
        "Scanopy-wireless01 EventId 08aa9805e843d6c6795f2a01e113de47 reached sequence 2899; netbox EventId 76de3509b456a3d5ecfbc00a57e84363 reached sequence 2916; "
        "socdev01 EventId 3e9d766fdd19b7572fa90228fa36299b reached sequence 2938.",
    )
    set_note(
        tables[5].cell(0, 0),
        "Relay and exclusions",
        "The server Wazuh relay reached checkpoint 3701 over TCP 514. Tpot01 is unsupported Debian 11; itarian01 and misp01 were unreachable and not enrolled. "
        "Those exclusions are outside the 26-host denominator, and relay status is not Wazuh decoder, rule, index or dashboard proof.",
    )
    set_row(
        tables[6],
        1,
        [
            "auditd baseline",
            "The per-host ProtectHome erratum was honored and its SHA-256 verified. MDWE/ProtectHome remained at each reviewed host's baseline; no hardening override was authorized.",
        ],
    )

    evidence_rows = [
        ["Deployment denominator", "Passed", "26/26 supported enrolled Linux endpoints; zero failures or unreachable targets inside the set."],
        ["Service and rules", "Passed", "0.9.6; four services enabled/active; six explicit-all rules; current-auditd-bound status and advancing health."],
        ["Attribution", "Passed", "Controlled reads produced Full attribution on every validated Linux endpoint."],
        ["Data preservation", "Passed", "Enrollment present and all 78/78 canary hashes preserved."],
        ["Management visibility", "Passed", "Inventory acknowledged 26 Linux 0.9.6 and two Windows 0.6.3; final event sequences recorded."],
    ]
    for index, values in enumerate(evidence_rows, 1):
        set_row(tables[7], index, values)
    set_note(
        tables[8].cell(0, 0),
        "Evidence ownership",
        "The rollout owner reported the live deployment and final independent read-only validation. This documentation pass reconciled those records and did not independently operate the endpoints or repeat the runtime tests.",
    )

    gate_rows = [
        ["Exact package", "Passed", "Immutable archive SHA-256 501a4ec589b3af6f90c2eca6b4a7762b3f0c45c80961b21872310c567b77f02e."],
        ["Supported Linux fleet", "Passed", "Deployment and final independent read-only validation accepted 26/26."],
        ["Canary and attribution", "Passed", "78/78 hashes preserved; Full attribution controlled reads on all 26 endpoints."],
        ["Excluded hosts", "Outside scope", "tpot01 unsupported; itarian01 and misp01 unreachable and not enrolled."],
        ["Repository-isolated host", "Passed", "Signed Debian 13 audit dependencies staged through Ansible for scanopy-wireless01; application remains unsigned."],
        ["Wazuh processing", "Pending", "TCP 514 relay checkpoint passed; decoder, rule, index and dashboard acceptance remain open."],
        ["Broader native recovery", "Pending", "Reboot and nonempty unrelated live-policy preservation remain open."],
        ["Platform and UI", "Pending", "ARM64 relay, broader operation/filesystem/mount/immutable-policy matrix and UI acceptance."],
        ["Production certification", "Pending", "Supported-fleet deployment acceptance does not certify the product for production."],
    ]
    for index, values in enumerate(gate_rows, 1):
        set_row(tables[9], index, values)

    set_note(
        tables[10].cell(0, 0),
        "Operational record",
        "Retain the exact 0.9.6 archive, per-host configuration evidence, validation results and immutable baseline erratum, SHA-256 59cd4a04c8ba61eb280a62b5cac2bdf22836a18737c3d8ed87b6f0c42bc675cd. "
        "Do not alter the archive or extracted manifests.",
    )
    operations = [
        "Treat the 26-host denominator as complete only for supported enrolled Linux endpoints covered by this deployment record.",
        "Keep tpot01, itarian01 and misp01 outside success metrics until support and enrollment requirements are independently resolved.",
        "Preserve automatic process-stop as disabled unless a separately approved response policy changes it.",
        "Complete production certification, Wazuh processing, reboot, broader native-policy, ARM64 relay and UI acceptance as separate release gates.",
    ]
    bullets = [p for p in document.paragraphs if p.style and p.style.name == "List Bullet"]
    for index, value in enumerate(operations, 4):
        bullets[index].text = value

    set_note(
        tables[11].cell(0, 0),
        "Decision",
        "The exact 0.9.6 archive is accepted for the completed 26-host supported enrolled Linux deployment. "
        "This closes the supported Linux fleet deployment gate recorded here. Production certification and the listed platform, recovery, UI and Wazuh-processing gates remain open.",
        True,
    )

    properties = document.core_properties
    properties.title = "Quiver seQuester Linux x64 0.9.6 Fleet Deployment Acceptance Addendum"
    properties.subject = "Supported enrolled Linux fleet deployment acceptance and remaining release gates"
    properties.author = "Quiver seQuester Documentation"
    properties.last_modified_by = "Quiver seQuester Documentation"
    properties.comments = ""

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Created {output}")


if __name__ == "__main__":
    main()
