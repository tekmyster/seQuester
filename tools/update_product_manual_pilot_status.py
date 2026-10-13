from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


OLD_BOUNDARY = (
    "IMPORTANT RELEASE BOUNDARY: all packages are unsigned development artifacts. Linux x64 0.9.1 "
    "build/hash/syntax checks and a narrow relay MDWE startup test passed, but full audit attribution did not. "
    "A later pilot initialized the relay/collector and advanced counters, then rolled back when auditd's "
    "ProtectHome namespace could not load a /home watch. No matching canary attribution or Wazuh acceptance "
    "was obtained. macOS C# Endpoint Security remains unimplemented. Do not claim production or platform parity."
)

NEW_BOUNDARY = (
    "IMPORTANT RELEASE BOUNDARY: all packages are unsigned development artifacts. Linux x64 0.9.1 "
    "build/hash/syntax checks and a narrow relay MDWE startup test passed, but full audit attribution did not. "
    "An earlier attempt produced partial transport/parser counter movement before rollback. The latest pilot kept "
    "the relay and collector processes stable, installed six exact rules, and preserved auditd "
    "MemoryDenyWriteExecute=yes; however, the agent rejected the relay connection repeatedly. Final health was "
    "Running=false, TransportConnected=false, RelayPid=0, InputBytes=0 and LastRelayFrameUtc=null, with heartbeat "
    "6 and 71 rejected connections. Rollout stopped before any canary read and restored healthy detection-only "
    "0.9.0. No matching canary attribution or Wazuh acceptance was obtained. macOS C# Endpoint Security remains "
    "unimplemented. Do not claim production or platform parity."
)

OLD_LIMITS = (
    "The recent work does not prove that a read completed a copy, classify every write as encryption, identify "
    "remote SMB users, provide macOS actor telemetry, create mobile endpoint parity, or make the product "
    "production-ready. Queues and relay buffers are finite. Linux audit rule coverage, kernel/auditd drops and "
    "service namespaces require independent checks. Global known-folder placement, per-instance central inventory, "
    "automatic correlation/scoring, learning mode, rich signer/parent attribution, advanced containment, signed "
    "releases, container deployment and complete provider integrations remain backlog items unless separately "
    "evidenced in the current source and QA ledger."
)

NEW_LIMITS = (
    "The recent work does not prove copy completion, encryption, remote SMB identity, macOS actor telemetry, "
    "mobile parity or production readiness. Queues are finite; Linux rule coverage, dropped audit records and "
    "service namespaces require independent checks. Global placement, central per-canary inventory, correlation, "
    "learning mode, richer attribution, containment, signing, containers and provider integrations remain backlog "
    "unless separately evidenced in source and the QA ledger."
)


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("output")
    args = parser.parse_args()

    document = Document(args.input)
    paragraph_updates = 0
    limits_updates = 0
    for paragraph in document.paragraphs:
        if paragraph.text == OLD_BOUNDARY:
            set_paragraph_text(paragraph, NEW_BOUNDARY)
            paragraph_updates += 1
        elif paragraph.text == OLD_LIMITS:
            set_paragraph_text(paragraph, NEW_LIMITS)
            limits_updates += 1

    table_updates = 0
    for table in document.tables:
        for row in table.rows:
            if row.cells[0].text == "Partial pipeline evidence":
                row.cells[0].text = "Latest pilot: connection rejected"
                row.cells[1].text = (
                    "Relay and collector processes remained stable; six exact rules were installed with auditd "
                    "MemoryDenyWriteExecute=yes. The agent reported Running=false, TransportConnected=false, "
                    "RelayPid=0, InputBytes=0, LastRelayFrameUtc=null, heartbeat 6 and 71 rejected connections."
                )
                row.cells[2].text = (
                    "Blocked. The pilot stopped before a canary read and rolled back to healthy detection-only "
                    "0.9.0. No live attribution or Wazuh acceptance was demonstrated."
                )
                table_updates += 1

    if paragraph_updates != 1 or limits_updates != 1 or table_updates != 1:
        raise RuntimeError(
            "Expected one boundary paragraph, one limits paragraph and one table update; "
            f"got boundary={paragraph_updates}, limits={limits_updates}, tables={table_updates}"
        )

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Updated latest pilot status: {output}")


if __name__ == "__main__":
    main()
