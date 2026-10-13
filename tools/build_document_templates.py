"""Build the fixed, synthetic DOCX/PDF payloads embedded in Ransom Free."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

root = Path(__file__).resolve().parents[1]
assets = root / "Assets"
assets.mkdir(exist_ok=True)

document = Document()
# Remove built-in title/paragraph border residue from the bundled base document.
for root_element in (document.styles.element, document.element):
    for border in list(root_element.xpath(".//w:pBdr")):
        border.getparent().remove(border)
section = document.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.85)
section.left_margin = section.right_margin = Inches(0.9)
for name in ("Normal", "Title", "Heading 1"):
    style = document.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.underline = False
document.styles["Normal"].font.size = Pt(11)
document.styles["Normal"].paragraph_format.space_after = Pt(10)
document.styles["Title"].font.size = Pt(24)
document.styles["Heading 1"].font.size = Pt(14)
document.core_properties.author = "Ransom Free"
document.core_properties.title = "Account recovery register"
document.add_paragraph("Account recovery register", "Title")
document.add_paragraph("This register records recovery references for retained accounts and operational records.")
document.add_paragraph("Recovery references", "Heading 1")
document.add_paragraph("Primary archive reference: DEMO-RECOVERY-001\nSecondary archive reference: DEMO-RECOVERY-002\nRecord owner: Example Operations")
document.add_paragraph("Retention notes", "Heading 1")
document.add_paragraph("Keep the recovery register with the associated archive inventory. Review the references when the inventory changes.")
document.add_paragraph("Synthetic security canary. These references do not grant access to any account or system.")
document.save(assets / "canary.docx")

styles = getSampleStyleSheet()
styles.add(ParagraphStyle("CanaryTitle", fontName="Helvetica-Bold", fontSize=24, leading=29, spaceAfter=20, textColor=colors.black))
styles.add(ParagraphStyle("CanaryBody", fontName="Helvetica", fontSize=11, leading=16, spaceAfter=12))
styles.add(ParagraphStyle("CanaryHeading", fontName="Helvetica-Bold", fontSize=14, leading=18, spaceBefore=12, spaceAfter=8))
pdf = SimpleDocTemplate(str(assets / "canary.pdf"), pagesize=letter,
    leftMargin=65, rightMargin=65, topMargin=62, bottomMargin=62,
    title="Confidential archive inventory", author="Ransom Free")
story = [Paragraph("Confidential archive inventory", styles["CanaryTitle"]),
    Paragraph("This inventory records references for retained operational archives.", styles["CanaryBody"]),
    Paragraph("Archive references", styles["CanaryHeading"]),
    Paragraph("DEMO-ARCHIVE-001 &nbsp; Operations records<br/>DEMO-ARCHIVE-002 &nbsp; Recovery register<br/>DEMO-ARCHIVE-003 &nbsp; Retention schedule", styles["CanaryBody"]),
    Paragraph("Retention notes", styles["CanaryHeading"]),
    Paragraph("Keep this inventory with the recovery register. Review the archive references when retained records change.", styles["CanaryBody"]),
    Spacer(1, 12),
    Paragraph("Synthetic security canary. No real credentials or personal data are included.", styles["CanaryBody"])]
pdf.build(story)
print("Created DOCX and PDF templates")
