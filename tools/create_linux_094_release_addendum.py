from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor

from create_linux_092_release_addendum import (
    RED,
    add_bullet,
    add_note,
    add_table,
    configure_section,
    configure_styles,
)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--logo", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    document = Document()
    configure_styles(document)
    configure_section(document.sections[0])

    header = document.sections[0].header.paragraphs[0]
    header.text = "Quiver seQuester  |  Linux 0.9.4 release status addendum"
    footer = document.sections[0].footer.paragraphs[0]
    for run in footer.runs:
        if "RF DOC 001-A" in run.text:
            run.text = run.text.replace("RF DOC 001-A", "RF DOC 001-C")

    logo_p = document.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo = logo_p.add_run().add_picture(args.logo, width=Inches(2.0))
    doc_pr = logo._inline.xpath("./wp:docPr")[0]
    doc_pr.set("descr", "Quiver seQuester shield and document logo")
    doc_pr.set("title", "Quiver seQuester logo")

    p = document.add_paragraph("Quiver product line", style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = document.add_paragraph("Linux x64 0.9.4 Release Status Addendum", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = document.add_paragraph("Current addendum to seQuester Product Manual, Revision 1.2", style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Published 13 September 2026  •  Activation blocked after native gate failure  •  Fleet rollout held")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(RED)
    document.add_paragraph()

    add_note(
        document,
        "Supersession",
        "This addendum supersedes RF DOC 001-B only for Linux x64 0.9.4 package and persistence-pilot status. "
        "Revision 1.2 and prior addenda remain immutable dated records. Use the packaged 0.9.4 persistence guide for upgrade, acceptance and rollback.",
    )
    add_note(
        document,
        "Acceptance boundary",
        "The immutable candidate passed package, planner, mocked service-lifecycle and isolated launcher checks, but failed the first native idempotence gate before mutation. "
        "All six exact rules were preserved; no auditd restart or canary read occurred. Do not activate this archive unchanged.",
        danger=True,
    )
    document.add_paragraph("Document control", style="Heading 1")
    add_table(
        document,
        ["Field", "Value"],
        [
            ["Document identifier", "RF DOC 001-C"],
            ["Supersedes", "RF DOC 001-B activation guidance for Linux x64 0.9.4 only"],
            ["Product", "Quiver seQuester"],
            ["Scope", "Linux x64 endpoint audit-attribution persistence candidate 0.9.4"],
            ["Status", "Activation blocked after native gate failure; replacement required; fleet rollout held"],
        ],
        [1.55, 5.8],
    )

    document.add_page_break()
    document.add_paragraph("1. Published package", style="Heading 1")
    add_table(
        document,
        ["Property", "Verified value"],
        [
            ["Release", "ransomfree-client-0.9.4-dev.20260913T175131Z-976d7320-linux-x64"],
            ["Archive", "ransomfree-client-0.9.4-dev.20260913T175131Z-976d7320-linux-x64.tar.gz"],
            ["SHA-256", "331af01449fbb6af4dcd9a135ac5edf3e03ff789d5bc35059898cc6e76976e11"],
            ["Size and inventory", "85,042,777 bytes; 478 files verified"],
            ["Endpoint / relay", "Endpoint 0.9.4; unchanged Native AOT relay 0.9.1, protocol 1"],
            ["Approval", "Unsigned development candidate; not production or fleet approved"],
        ],
        [1.55, 5.8],
    )

    document.add_paragraph("2. Persistence correction", style="Heading 1")
    add_bullet(document, "The .NET rule planner computes only missing approved seQuester rules from current audit inventory.")
    add_bullet(document, "The reconciler adds missing exact rules without a global clear, unrelated-rule deletion or blind retry loop.")
    add_bullet(document, "The service dependency binds reconciliation to auditd start/restart transactions; an upgrade must re-enable the changed loader unit.")
    add_bullet(document, "Health must confirm that rules belong to the current auditd root/start instance, rather than treating a connected transport as complete protection.")
    add_note(
        document,
        "0.9.3 boundary",
        "The verified 0.9.3 pilot restored transport after auditd restart but lost all six scoped rules until an administrator restarted the loader. "
        "0.9.4 is designed to correct that persistence defect; fixture evidence is not a substitute for the live gate.",
        danger=True,
    )
    add_note(
        document,
        "0.9.4 live failure",
        "The initial native idempotence check reported Expected=6, Added=0, Success=false and Pid=0, then raised InvalidDataException and terminated with SIGABRT. "
        "The leading diagnosis is an overly strict parser for a standard three-token auditctl -s loginuid_immutable line; read-only confirmation is still pending. No rule mutation occurred.",
        danger=True,
    )

    document.add_paragraph("3. Preserved security contract", style="Heading 1")
    add_table(
        document,
        ["Control", "Required state"],
        [
            ["auditd", "Keep MemoryDenyWriteExecute=yes and ProtectHome=yes; do not weaken the host service."],
            ["Plugin", "Zero-argument root-owned launch-relay.sh; no JIT plugin and no capability additions."],
            ["Relay / collector", "Root-owned, capability-free, NoNewPrivileges-enabled transport identity contract remains unchanged."],
            ["Rules", "Six administrator-approved scoped rules plus all unrelated baseline rules; no global deletion."],
            ["Data", "Preserve enrollment, queues, cursor, protected inbox and every existing canary during upgrade or rollback."],
        ],
        [1.55, 5.8],
    )

    document.add_page_break()
    document.add_paragraph("4. Verified pre-pilot evidence", style="Heading 1")
    add_table(
        document,
        ["Evidence", "Result", "Limit"],
        [
            ["Package integrity", "Passed", "Complete archive/payload inventory, ELF/Bash/static contracts and tested-source hash gates; 478 files."],
            ["Rule planner", "Passed", "19 in-memory production-planner cases; no native auditctl I/O."],
            ["Dependency lifecycle", "Passed", "Real user-systemd lifecycle over mock audit inventory; one reconciliation per transaction, unrelated baseline preserved."],
            ["Reconciliation behavior", "Passed", "Two missing rules added after initial load, two dummy restarts and stop/start; explicit loader restart added zero."],
            ["Launcher transport", "Passed", "Exact launcher with inherited MDWE, empty environment, zero capabilities and NoNewPrivileges; not actual auditd/system-service startup."],
        ],
        [1.55, 1.0, 4.8],
    )
    add_note(
        document,
        "Evidence files",
        "The package includes tools/audit-rule-evidence-0.9.4.json and tools/audit-launcher-evidence-0.9.4.json. "
        "The agent executable and DLL match the frozen native fixture binaries byte for byte.",
    )

    document.add_paragraph("5. Live acceptance gates", style="Heading 1")
    add_table(
        document,
        ["Gate", "Status", "Required proof"],
        [
            ["Initial native idempotence", "Failed", "Failed before mutation: Expected=6, Added=0, Success=false, Pid=0; InvalidDataException followed by SIGABRT."],
            ["Rule preservation", "Passed", "All six exact rules remained; no auditd restart or canary read occurred."],
            ["Native auditctl and sandbox", "Blocked", "Confirm and correct native auditctl status parsing, then publish a replacement candidate before retry."],
            ["Automatic auditd recovery", "Pending", "Two actual auditd restarts with six exact rules restored automatically, baseline preserved and no manual loader rescue."],
            ["Idempotence and health", "Pending", "Explicit reload adds zero; health binds to the current auditd instance and reports any gap."],
            ["Live attribution/delivery", "Pending", "Controlled unchanged-canary read with Full attribution and the same EventId retained by management."],
            ["Wazuh", "Pending", "Wire arrival is not decoder, rule, index or dashboard acceptance; prove each required layer separately."],
            ["Recovery and UI", "Pending", "Reboot/mount/immutable-mode recovery, rollback and accurate operator status presentation."],
            ["Fleet rollout", "Held", "All required pilot evidence reviewed and accepted by the rollout owner."],
        ],
        [2.0, 1.05, 4.3],
    )

    document.add_page_break()
    document.add_paragraph("6. Operational hold and recovery", style="Heading 1")
    add_note(
        document,
        "Authoritative procedure",
        "The packaged packaging/linux-client/AUDIT_PERSISTENCE_0_9_4.md records the candidate's controlled procedure, but the failed native gate now supersedes its activation path. "
        "Use it only for evidence interpretation and the explicitly reviewed recovery steps. Do not activate this archive unchanged or substitute the historical AUDIT_PILOT.md procedure.",
    )
    add_bullet(document, "Hold the current 0.9.4 archive for diagnosis. Do not use it for a new activation or retry unchanged.")
    add_bullet(document, "Capture existing audit rules and service state first. Preserve unrelated rules and the existing endpoint identity, queues, cursor, inbox and canaries.")
    add_bullet(document, "Stop immediately on a failed gate. Preserve evidence and follow the packaged recovery or 0.9.3 rollback path; rollback returns to a manually maintained pilot, not fleet-ready protection.")
    add_bullet(document, "Do not claim lossless coverage, production readiness or Wazuh indexing from package, fixture or wire-capture evidence.")

    document.add_paragraph("7. Current references", style="Heading 1")
    for item in (
        "ATTRIBUTION_RELEASE.md",
        "UNIX_CLIENTS_RELEASE.md",
        "packaging/linux-client/AUDIT_PERSISTENCE_0_9_4.md",
        "CLIENT_FEATURE_MATRIX.md",
        "CURRENT_PRODUCT_STATUS.md",
        "DOCUMENTATION_CHANGE_LEDGER.md",
        "QA.md",
    ):
        add_bullet(document, item)

    document.add_paragraph("8. Release decision", style="Heading 1")
    add_note(
        document,
        "Decision",
        "Do not activate the current 0.9.4 archive unchanged. Preserve it as immutable failure evidence. Confirm the parser diagnosis, correct the defect, publish a replacement immutable candidate and restart acceptance from the failed native gate. Production and fleet rollout remain blocked.",
        danger=True,
    )

    properties = document.core_properties
    properties.title = "Quiver seQuester Linux x64 0.9.4 Release Status Addendum"
    properties.subject = "Immutable Linux x64 0.9.4 development candidate and pending pilot gates"
    properties.author = "Quiver seQuester Documentation"
    properties.last_modified_by = "Quiver seQuester Documentation"
    properties.keywords = "Quiver, seQuester, Linux, auditd, release status, persistence"
    properties.comments = ""

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Created {output}")


if __name__ == "__main__":
    main()
