from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def replace_paragraph_text(paragraph, old: str, new: str) -> bool:
    if old not in paragraph.text:
        return False
    replacement = paragraph.text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = replacement
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replacement)
    return True


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, value, end):
        run._r.append(node)


def set_footer(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = 2
    paragraph.add_run("RF DOC 001  •  Revision 1.2  |  Page ")
    page_run = paragraph.add_run()
    add_field(page_run, " PAGE ")
    paragraph.add_run(" of ")
    total_run = paragraph.add_run()
    add_field(total_run, " NUMPAGES ")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("output")
    args = parser.parse_args()

    document = Document(args.input)
    cover_updates = 0
    for paragraph in document.paragraphs:
        if replace_paragraph_text(paragraph, "Revision 1.1", "Revision 1.2"):
            cover_updates += 1

    seen_parts: set[str] = set()
    footer_updates = 0
    for section in document.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            part_name = str(footer.part.partname)
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            if not footer.paragraphs:
                paragraph = footer.add_paragraph()
            else:
                paragraph = footer.paragraphs[0]
            set_footer(paragraph)
            footer_updates += 1

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Updated {cover_updates} cover paragraph(s) and {footer_updates} footer part(s): {output}")


if __name__ == "__main__":
    main()
