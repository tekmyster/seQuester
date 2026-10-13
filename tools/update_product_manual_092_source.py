from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


OLD_SNAPSHOT = (
    "Documentation snapshot: 13 September 2026. Linux x64 0.9.1 adds a C# Native AOT auditd relay and "
    "separate managed collector. The published archive is not deployable unchanged because its empty args "
    "directive is rejected by auditd 3.0.7 and its shell rule draft is not persistent augenrules syntax. Linux "
    "ARM64 remains on the earlier client baseline without a matching 0.9.1 relay. macOS packages remain unchanged "
    "FSEvents baselines."
)

NEW_SNAPSHOT = (
    "Documentation snapshot: 13 September 2026. The published Linux x64 0.9.1 archive is not deployable "
    "unchanged because auditd 3.0.7 rejects its empty args directive and its shell rule draft is not persistent "
    "augenrules syntax. Unpublished 0.9.2 source adds bounded transport diagnostics, a restricted direct rule-file "
    "generator and a separate protected loader; it is not a released or pilot-proven fix. Linux ARM64 remains on "
    "the earlier baseline without a matching Native AOT relay. macOS packages remain FSEvents baselines."
)

OLD_ROLLBACK = (
    "The rollout-owned pilot returned to healthy detection-only 0.9.0 after rollback, with auditd "
    "MemoryDenyWriteExecute still enabled. The proposed post-auditd root oneshot loader and seQuester-only rules "
    "file outside /etc/audit/rules.d are pilot deployment work, not shipped product support. Acceptance must cover "
    "service ordering, reboot, failure visibility, exact removal, duplicate prevention and preservation of "
    "unrelated audit/Wazuh rules."
)

NEW_ROLLBACK = (
    "The pilot returned to healthy detection-only 0.9.0 after rollback, with auditd MemoryDenyWriteExecute still "
    "enabled. Unpublished 0.9.2 source adds bounded stage/category/errno/PID transport diagnostics and a separate "
    "sequester-audit-rules.service that reads an administrator-reviewed root-owned seQuester-only file outside "
    "/etc/audit/rules.d. It does not log payloads, paths, command lines or exception text and does not promise "
    "automatic reconciliation or duplicate-free reloads. Packaging and acceptance remain pending."
)


def set_text(paragraph, value: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("output")
    args = parser.parse_args()

    document = Document(args.input)
    snapshot_updates = 0
    rollback_updates = 0
    for paragraph in document.paragraphs:
        if paragraph.text == OLD_SNAPSHOT:
            set_text(paragraph, NEW_SNAPSHOT)
            snapshot_updates += 1
        elif paragraph.text == OLD_ROLLBACK:
            set_text(paragraph, NEW_ROLLBACK)
            rollback_updates += 1

    command_updates = 0
    for table in document.tables:
        if not table.rows or table.rows[0].cells[0].text != "Command":
            continue
        if not any(row.cells[0].text == "--audit-rules" for row in table.rows):
            continue
        if any(row.cells[0].text == "--audit-rules-file" for row in table.rows):
            continue
        row = table.add_row()
        row.cells[0].text = "--audit-rules-file"
        row.cells[1].text = (
            "Unpublished 0.9.2 source: emits direct auditctl rule-file syntax only for paths limited to ASCII "
            "letters, digits, slash, underscore, dash and dot. Other names must use reviewed shell-draft handling."
        )
        command_updates += 1

    if (snapshot_updates, rollback_updates, command_updates) != (1, 1, 1):
        raise RuntimeError(
            f"Expected one snapshot, rollback and command update; got "
            f"{snapshot_updates}, {rollback_updates}, {command_updates}"
        )

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Updated unpublished 0.9.2 source status: {output}")


if __name__ == "__main__":
    main()
