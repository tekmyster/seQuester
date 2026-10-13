from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from create_linux_092_release_addendum import GRAY, RED, add_field
from create_linux_095_release_addendum import set_note, set_row


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    document = Document(args.template)
    document.sections[0].header.paragraphs[0].text = "Quiver seQuester  |  Linux 0.9.5 scoped pilot acceptance"
    footer = document.sections[0].footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("RF DOC 001-E  •  13 September 2026  |  Page ")
    add_field(footer.add_run(), " PAGE ", "1")
    footer.add_run(" of ")
    add_field(footer.add_run(), " NUMPAGES ", "4")
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)

    replacements = {
        "Linux x64 0.9.5 Release Status Addendum": "Linux x64 0.9.5 Scoped Pilot Acceptance Addendum",
        "2. Persistence correction": "2. Accepted persistence behavior",
        "4. Verified pre-pilot evidence": "4. Scoped native acceptance evidence",
        "5. Live acceptance gates": "5. Acceptance results and remaining gates",
        "6. Controlled pilot and rollback": "6. Controlled rollout boundary",
    }
    for paragraph in document.paragraphs:
        if paragraph.text in replacements:
            paragraph.text = replacements[paragraph.text]
        if "Corrected unsigned development candidate" in paragraph.text:
            paragraph.text = "Published 13 September 2026  •  Scoped native pilot accepted  •  Controlled rollout in progress"
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(RED)

    tables = document.tables
    set_note(
        tables[0].cell(0, 0),
        "Supersession",
        "This addendum supersedes RF DOC 001-D only for the reported Linux x64 0.9.5 plex01 pilot result. "
        "RF DOC 001-D remains the immutable publication-time pending record, and RF DOC 001-C remains the immutable blocked-0.9.4 record.",
    )
    set_note(
        tables[1].cell(0, 0),
        "Acceptance boundary",
        "The rollout owner reports that exact 0.9.5 native reconciliation, two auditd restart restorations, zero-add loader restart, service controls, Full attribution and management/Wazuh wire delivery passed on plex01. "
        "This is scoped pilot acceptance, not completed fleet rollout, production certification or Wazuh decoder/index/dashboard acceptance.",
        True,
    )
    set_row(tables[2], 1, ["Document identifier", "RF DOC 001-E"])
    set_row(tables[2], 2, ["Supersedes", "RF DOC 001-D pilot-status guidance only"])
    set_row(tables[2], 4, ["Scope", "Linux x64 0.9.5 scoped native plex01 pilot acceptance"])
    set_row(tables[2], 5, ["Status", "Scoped pilot accepted; controlled rollout in progress; production certification pending"])
    set_row(tables[3], 6, ["Approval", "Scoped plex01 pilot accepted; controlled rollout only; not production or completed fleet approved"])

    behavior = [
        "Initial reconciliation found six expected and six already present rules and added zero.",
        "Two actual distribution auditd restarts each triggered exactly one new loader InvocationID and RunId and automatically restored six unique rules without manual rescue.",
        "One explicit loader-only restart created one new invocation and added zero rules.",
        "Root health passed after settling and bound protection to the current auditd instance.",
    ]
    bullets = [p for p in document.paragraphs if p.style and p.style.name == "List Bullet"]
    for index, value in enumerate(behavior):
        bullets[index].text = value
    set_note(
        tables[4].cell(0, 0),
        "0.9.3 boundary",
        "The verified 0.9.3 pilot restored transport after auditd restart but lost all six scoped rules until an administrator restarted the loader. "
        "The scoped 0.9.5 pilot accepted automatic restoration across two actual auditd restarts; broader rollout and production gates remain separate.",
        True,
    )

    evidence_rows = [
        ["Package identity", "Passed", "Exact 0.9.5 archive and SHA-256 accepted on plex01; package remained immutable."],
        ["Native reconciliation", "Passed", "Expected 6, already present 6, added 0; protected status and current-auditd health passed after settling."],
        ["Auditd restart recovery", "Passed", "Two actual restarts; exactly one loader run each; six unique rules restored automatically; no manual rescue."],
        ["Service controls", "Passed", "Collector and auditd-child relay root:root, capability sets zero, NNP1; MDWE and ProtectHome remained enabled."],
        ["Attributed delivery", "Passed", "Unchanged payroll read produced Full attribution; management checkpoint 1054 and exact JSON wire arrival at Wazuh TCP 514."],
    ]
    for index, values in enumerate(evidence_rows, 1):
        set_row(tables[7], index, values)
    set_note(
        tables[8].cell(0, 0),
        "Rollout owner evidence",
        "EventId ee21bc71a11788ac4c0f562f3ab72752; /usr/bin/cat PID 206298, PPID 206284, UID/eUID 0; management sequence and relay checkpoint 1054; exact JSON captured at the configured Wazuh TCP 514 receiver from CPTNYC_DC01. "
        "The payroll canary SHA-256, enrollment and all three canary hashes were preserved.",
    )

    gate_rows = [
        ["Exact package", "Passed", "SHA-256 41f8e54305efa51e49833e375060dd53f260f9738e37fdd347ad5bf5b434a6df."],
        ["Native reconciliation", "Passed", "Initial six-rule inventory reconciled idempotently with zero additions."],
        ["Automatic auditd recovery", "Passed", "Two actual auditd restarts automatically restored six unique rules with one loader run each."],
        ["Loader idempotence", "Passed", "Explicit loader-only restart produced one invocation and added zero."],
        ["Full attribution/delivery", "Passed", "Controlled unchanged-canary read reached management with exact process/user fields and EventId."],
        ["Wazuh wire arrival", "Passed", "Exact JSON reached TCP 514; decoder, rule, index and dashboard acceptance remain pending."],
        ["Broader native matrix", "Pending", "Nonempty unrelated live baseline, ARM64, broader users/operations/filesystems, mounts and immutable-policy modes."],
        ["Reboot and UI", "Pending", "Reboot recovery and accurate desktop/management status presentation."],
        ["Fleet and production", "In progress", "Controlled rollout is proceeding under existing authorization; completion and production certification are not reported."],
    ]
    for index, values in enumerate(gate_rows, 1):
        set_row(tables[9], index, values)

    set_note(
        tables[10].cell(0, 0),
        "Authoritative procedure",
        "Continue to use packaging/linux-client/AUDIT_PERSISTENCE_0_9_5.md for controlled deployment and rollback. "
        "Do not deploy blocked 0.9.4 or treat this scoped pilot as broad platform acceptance.",
    )
    rollout_bullets = [
        "Controlled Linux rollout may proceed only under the existing authorization and operational safeguards; completion has not been reported.",
        "Preserve unrelated rules, endpoint identity, queues, cursor, inbox and canaries; retain per-host evidence and exact package identity.",
        "Stop expansion on any new blocker or failed gate, preserve evidence, and follow the packaged 0.9.5 recovery or reviewed 0.9.3 rollback path.",
        "Do not claim production certification, completed fleet coverage, lossless restart behavior or Wazuh indexing from this evidence.",
    ]
    bullets = [p for p in document.paragraphs if p.style and p.style.name == "List Bullet"]
    for index, value in enumerate(rollout_bullets, 4):
        bullets[index].text = value

    set_note(
        tables[11].cell(0, 0),
        "Decision",
        "The exact Linux x64 0.9.5 package passed the scoped plex01 native pilot and may continue through the separately authorized controlled rollout. "
        "Fleet completion, production certification, ARM64, reboot, broader coverage, UI acceptance and Wazuh processing remain open.",
        True,
    )

    properties = document.core_properties
    properties.title = "Quiver seQuester Linux x64 0.9.5 Scoped Pilot Acceptance Addendum"
    properties.subject = "Rollout-owner-reported Linux x64 0.9.5 scoped native pilot acceptance and remaining gates"
    properties.author = "Quiver seQuester Documentation"
    properties.last_modified_by = "Quiver seQuester Documentation"
    properties.comments = ""

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Created {output}")


if __name__ == "__main__":
    main()
