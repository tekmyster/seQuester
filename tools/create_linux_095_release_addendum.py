from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from create_linux_092_release_addendum import GRAY, NAVY, RED, add_field, set_cell_text


def set_note(cell, label: str, text: str, danger: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(f"{label}: ")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(RED if danger else NAVY)
    paragraph.add_run(text)


def set_row(table, index: int, values: list[str]) -> None:
    for column, value in enumerate(values):
        set_cell_text(table.rows[index].cells[column], value)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    document = Document(args.template)
    document.sections[0].header.paragraphs[0].text = "Quiver seQuester  |  Linux 0.9.5 release status addendum"
    footer = document.sections[0].footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("RF DOC 001-D  •  13 September 2026  |  Page ")
    add_field(footer.add_run(), " PAGE ", "1")
    footer.add_run(" of ")
    add_field(footer.add_run(), " NUMPAGES ", "4")
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)

    replacements = {
        "Linux x64 0.9.4 Release Status Addendum": "Linux x64 0.9.5 Release Status Addendum",
        "6. Operational hold and recovery": "6. Controlled pilot and rollback",
    }
    for paragraph in document.paragraphs:
        if paragraph.text in replacements:
            paragraph.text = replacements[paragraph.text]
        if "Activation blocked after native gate failure" in paragraph.text:
            paragraph.text = "Published 13 September 2026  •  Corrected unsigned development candidate  •  Fleet rollout held"
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(RED)

    tables = document.tables
    set_note(
        tables[0].cell(0, 0),
        "Supersession",
        "This addendum supersedes RF DOC 001-C only for Linux x64 0.9.5 package and parser-correction status. "
        "RF DOC 001-C remains an immutable record of the blocked 0.9.4 candidate. Use the packaged 0.9.5 persistence guide for the bounded pilot.",
    )
    set_note(
        tables[1].cell(0, 0),
        "Acceptance boundary",
        "The immutable 0.9.5 candidate, package integrity, bounded parser, managed failure exit, planner, mocked lifecycle and isolated launcher transport are verified. "
        "Successful native reconciliation, actual auditd restart/reboot, live attribution/delivery and UI acceptance remain pending.",
        True,
    )
    set_row(tables[2], 1, ["Document identifier", "RF DOC 001-D"])
    set_row(tables[2], 2, ["Supersedes", "RF DOC 001-C activation guidance for Linux x64 0.9.5 only"])
    set_row(tables[2], 4, ["Scope", "Linux x64 endpoint audit-attribution parser correction 0.9.5"])
    set_row(tables[2], 5, ["Status", "Corrected development candidate; live persistence pilot pending; fleet rollout held"])

    package_rows = [
        ["Release", "ransomfree-client-0.9.5-dev.20260913T181115Z-52d56843-linux-x64"],
        ["Archive", "ransomfree-client-0.9.5-dev.20260913T181115Z-52d56843-linux-x64.tar.gz"],
        ["SHA-256", "41f8e54305efa51e49833e375060dd53f260f9738e37fdd347ad5bf5b434a6df"],
        ["Size and inventory", "85,047,394 bytes; 480 files verified"],
        ["Endpoint / relay", "Endpoint 0.9.5; unchanged Native AOT relay 0.9.1, protocol 1"],
        ["Approval", "Unsigned development candidate; existing bounded pilot only; not production or fleet approved"],
    ]
    for index, values in enumerate(package_rows, 1):
        set_row(tables[3], index, values)

    set_note(
        tables[5].cell(0, 0),
        "0.9.4 failure corrected",
        "The 0.9.4 parser rejected the valid audit status line loginuid_immutable 0 unlocked and an unhandled failure ended in SIGABRT. "
        "0.9.5 accepts that bounded optional annotation while still rejecting malformed, duplicate or oversized identity output. Operational failures now return exit code 1 and persist FailureStage without a stack trace or abort.",
        True,
    )

    evidence_rows = [
        ["Package integrity", "Passed", "Complete archive/payload, ELF/Bash/static and exact-source/fixture checks; 480 files."],
        ["Status parser and exit", "Passed", "17 local/native cases using exact observed auditd 3.0.7 output; malformed/duplicate identities rejected."],
        ["Native error path", "Passed", "Actual agent missing-reviewed-file case returned 1 and persisted failed DesiredRules stage/counts; no abort, stack or auditctl call."],
        ["Rule planner/lifecycle", "Passed", "Nineteen planner cases rerun; prior user-systemd/mock-inventory lifecycle proof retained within its original limits."],
        ["Launcher transport", "Passed", "Exact 0.9.5 launcher under inherited MDWE and empty environment; root:root, capability-free, NNP1, 88 bytes, heartbeat 3, failures 0."],
    ]
    for index, values in enumerate(evidence_rows, 1):
        set_row(tables[7], index, values)
    set_note(
        tables[8].cell(0, 0),
        "Evidence files",
        "The package includes tools/audit-status-evidence-0.9.5.json and tools/audit-launcher-evidence-0.9.5.json. "
        "The agent executable and DLL match the frozen native fixture binaries byte for byte. Successful native reconciliation was not exercised.",
    )

    gate_rows = [
        ["Parser and managed exit", "Passed", "Seventeen local/native cases plus actual pre-auditctl failure returned exit 1 with protected FailureStage."],
        ["Package/fixture identity", "Passed", "Agent executable and DLL match frozen native fixture inputs; complete 480-file verification."],
        ["Native reconciliation", "Pending", "Successful auditctl inventory/add operations under the packaged system-service sandbox."],
        ["Automatic auditd recovery", "Pending", "Two actual auditd restarts restore six exact rules automatically, preserve baseline and require no manual rescue."],
        ["Idempotence and health", "Pending", "Explicit reload adds zero; health binds to the current auditd instance and reports any gap."],
        ["Live attribution/delivery", "Pending", "Controlled unchanged-canary read with Full attribution and the same EventId retained by management."],
        ["Wazuh", "Pending", "Prove required decoder, rule, index and dashboard layers; wire arrival alone is insufficient."],
        ["Recovery and UI", "Pending", "Reboot/mount/immutable-mode recovery, rollback and accurate operator status presentation."],
        ["Fleet rollout", "Held", "All required pilot evidence reviewed and accepted by the rollout owner."],
    ]
    for index, values in enumerate(gate_rows, 1):
        set_row(tables[9], index, values)

    set_note(
        tables[10].cell(0, 0),
        "Authoritative procedure",
        "Use packaging/linux-client/AUDIT_PERSISTENCE_0_9_5.md from the exact 0.9.5 candidate. It defines protected maintenance, inventory capture, acceptance and rollback. "
        "Do not deploy blocked 0.9.4 unchanged or substitute historical 0.9.4/0.9.3 procedures.",
    )
    bullet_map = {
        "Hold the current 0.9.4 archive for diagnosis. Do not use it for a new activation or retry unchanged.":
            "Use 0.9.5 only for the already approved bounded pilot. Confirm its exact archive hash and inventory before changing the selector.",
        "Capture existing audit rules and service state first. Preserve unrelated rules and the existing endpoint identity, queues, cursor, inbox and canaries.":
            "Capture existing audit rules and service state first. Preserve unrelated rules, endpoint identity, queues, cursor, inbox and canaries.",
        "Stop immediately on a failed gate. Preserve evidence and follow the packaged recovery or 0.9.3 rollback path; rollback returns to a manually maintained pilot, not fleet-ready protection.":
            "Stop immediately on a failed gate. Preserve evidence and follow the 0.9.5 packaged recovery or 0.9.3 rollback path; rollback remains manually maintained.",
    }
    for paragraph in document.paragraphs:
        if paragraph.text in bullet_map:
            paragraph.text = bullet_map[paragraph.text]
        if paragraph.text == "packaging/linux-client/AUDIT_PERSISTENCE_0_9_4.md":
            paragraph.text = "packaging/linux-client/AUDIT_PERSISTENCE_0_9_5.md"

    set_note(
        tables[11].cell(0, 0),
        "Decision",
        "The 0.9.5 archive is suitable only for the existing bounded persistence pilot. It is not approved for production, fleet rollout, successful native reconciliation or automatic restart persistence until every pending live gate passes.",
        True,
    )

    properties = document.core_properties
    properties.title = "Quiver seQuester Linux x64 0.9.5 Release Status Addendum"
    properties.subject = "Corrected Linux x64 0.9.5 development candidate and pending pilot gates"
    properties.author = "Quiver seQuester Documentation"
    properties.last_modified_by = "Quiver seQuester Documentation"
    properties.comments = ""

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Created {output}")


if __name__ == "__main__":
    main()
