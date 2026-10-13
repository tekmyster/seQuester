from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from create_linux_092_release_addendum import GRAY, RED, add_field
from create_linux_095_release_addendum import set_note, set_row


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    document = Document(args.template)
    document.sections[0].header.paragraphs[0].text = "Quiver seQuester  |  Linux 0.9.6 compatibility candidate"
    footer = document.sections[0].footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("RF DOC 001-F  •  13 September 2026  |  Page ")
    add_field(footer.add_run(), " PAGE ", "1")
    footer.add_run(" of ")
    add_field(footer.add_run(), " NUMPAGES ", "4")
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)

    replacements = {
        "Linux x64 0.9.5 Release Status Addendum": "Linux x64 0.9.6 Compatibility Candidate Addendum",
        "2. Persistence correction": "2. Explicit syscall selector correction",
    }
    for paragraph in document.paragraphs:
        if paragraph.text in replacements:
            paragraph.text = replacements[paragraph.text]
        if "Corrected unsigned development candidate" in paragraph.text:
            paragraph.text = "Published 13 September 2026  •  Replacement compatibility pilots pending  •  Fleet expansion blocked"
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(RED)

    tables = document.tables
    set_note(
        tables[0].cell(0, 0),
        "Supersession",
        "This addendum supersedes RF DOC 001-E only for current Linux x64 0.9.6 compatibility-candidate status. "
        "RF DOC 001-E remains immutable single-host plex01 evidence, and prior addenda retain their historical scopes.",
    )
    set_note(
        tables[1].cell(0, 0),
        "Acceptance boundary",
        "The immutable package and dual-libaudit in-memory mask fixtures are verified. The fixtures performed no native kernel audit operations. "
        "Actual 0.9.6 Ubuntu and Debian reconciliation, two auditd restarts, controlled read/delivery, reboot and UI acceptance remain pending; fleet expansion is blocked.",
        True,
    )
    set_row(tables[2], 1, ["Document identifier", "RF DOC 001-F"])
    set_row(tables[2], 2, ["Supersedes", "RF DOC 001-E current-candidate guidance only"])
    set_row(tables[2], 4, ["Scope", "Linux x64 0.9.6 explicit syscall-selector compatibility candidate"])
    set_row(tables[2], 5, ["Status", "Unsigned development candidate; replacement pilots pending; fleet expansion blocked"])

    package_rows = [
        ["Release", "ransomfree-client-0.9.6-dev.20260913T184639Z-08804ca0-linux-x64"],
        ["Archive", "ransomfree-client-0.9.6-dev.20260913T184639Z-08804ca0-linux-x64.tar.gz"],
        ["SHA-256", "501a4ec589b3af6f90c2eca6b4a7762b3f0c45c80961b21872310c567b77f02e"],
        ["Size and inventory", "85,053,726 bytes; 481 files verified"],
        ["Endpoint / relay", "Endpoint 0.9.6; unchanged Native AOT relay 0.9.1, protocol 1"],
        ["Approval", "Replacement compatibility pilots only; not production or fleet approved"],
    ]
    for index, values in enumerate(package_rows, 1):
        set_row(tables[3], index, values)

    correction = [
        "The approved desired rule retains explicit -S all in the actual ManagedAuditRule add arguments.",
        "Native in-memory libaudit 3.0.7 and 4.0.2 fixtures retained the full 2048-bit mask for base64 and base32 rule forms.",
        "The legacy implicit command reproduced the exact nginx01 46-syscall list on 4.0.2; 3.0.7 left the mask unset before its default fallback.",
        "The correction does not accept arbitrary or preexisting expanded syscall lists as equivalent; those remain fail-closed conflicts requiring operator review.",
    ]
    bullets = [p for p in document.paragraphs if p.style and p.style.name == "List Bullet"]
    for index, value in enumerate(correction):
        bullets[index].text = value
    set_note(
        tables[4].cell(0, 0),
        "0.9.5 evidence boundary",
        "Plex01 Ubuntu/auditctl 3.0.7 remains valid single-host acceptance. Nginx01 Debian/auditctl 4.0.2 exposed the dropped-selector defect after one partial rule; scoped cleanup restored original 0.8.0 with no remaining partial rule or integration.",
        True,
    )
    set_note(
        tables[5].cell(0, 0),
        "Correction scope",
        "0.9.6 preserves the approved argument that 0.9.5 dropped. It does not weaken rule identity or normalize arbitrary expanded lists into trusted equivalence.",
        True,
    )
    set_row(
        tables[6],
        1,
        [
            "auditd baseline",
            "Preserve each host's effective baseline: MDWE=yes on reviewed hosts; ProtectHome=yes on Ubuntu plex01 and no on Debian nginx01/backup01. No hardening override is authorized or required.",
        ],
    )

    evidence_rows = [
        ["Package integrity", "Passed", "Complete immutable archive/payload verified; 481 files; agent binaries match the frozen native fixture."],
        ["Libaudit mask fixtures", "Passed", "3.0.7 and 4.0.2 retained explicit -S all as 2048 bits for base64/base32; no kernel audit operations."],
        ["Legacy reproduction", "Passed", "Implicit command reproduced nginx01's exact 46-syscall 4.0.2 list and the 3.0.7 unset-mask behavior."],
        ["Planner and parser", "Passed", "All 23 planner and 17 status/exit cases passed on both native fixture versions."],
        ["Launcher and lifecycle", "Passed", "Current MDWE/capability-free transport, normal error path and updated-planner user-systemd lifecycle passed."],
    ]
    for index, values in enumerate(evidence_rows, 1):
        set_row(tables[7], index, values)
    set_note(
        tables[8].cell(0, 0),
        "Evidence limit",
        "These fixtures validate rule construction, parsing, launcher and mocked lifecycle behavior. They did not add, enumerate or remove rules in a native kernel audit subsystem and do not prove replacement-host reconciliation or live delivery.",
    )

    gate_rows = [
        ["Exact package", "Passed", "SHA-256 501a4ec589b3af6f90c2eca6b4a7762b3f0c45c80961b21872310c567b77f02e."],
        ["Dual-libaudit fixtures", "Passed", "Explicit selector retained on 3.0.7 and 4.0.2; legacy behavior reproduced."],
        ["Ubuntu replacement pilot", "Pending", "Existing six-rule host must reconcile expected 6/already present 6/added 0."],
        ["Debian clean pilot", "Pending", "Restored host must reconcile expected 6/already present 0/added 6 without conflict or partial accounting."],
        ["Automatic restart recovery", "Pending", "Two actual auditd restarts per replacement host restore six rules without manual rescue."],
        ["Attributed delivery", "Pending", "Controlled unchanged-canary read reaches management with exact EventId/process/user attribution."],
        ["Wazuh", "Pending", "Wire arrival and decoder/rule/index/dashboard acceptance require separate evidence."],
        ["Broader recovery and UI", "Pending", "Reboot, ARM64, broader coverage, mounts/immutable modes and operator UI acceptance."],
        ["Fleet expansion", "Blocked", "Resume only after both replacement compatibility pilots pass and rollout owner approves expansion."],
    ]
    for index, values in enumerate(gate_rows, 1):
        set_row(tables[9], index, values)

    set_note(
        tables[10].cell(0, 0),
        "Authoritative procedure",
        "Use the corrected packaging/linux-client/AUDIT_PERSISTENCE_0_9_6.md from the exact candidate together with the immutable standalone baseline erratum (SHA-256 59cd4a04c8ba61eb280a62b5cac2bdf22836a18737c3d8ed87b6f0c42bc675cd). "
        "Preserve each host's effective ProtectHome state; do not alter extracted package manifests.",
    )
    rollout = [
        "Run only the approved replacement compatibility pilots; fleet expansion remains blocked.",
        "Capture desired, issued and live rule forms plus exact auditctl/libaudit versions before and after each reconciliation.",
        "Treat any preexisting expanded syscall list as a fail-closed conflict requiring explicit operator review; do not auto-trust or broadly clear it.",
        "Preserve endpoint identity, queues, cursor, inbox and canaries; follow the package's host-specific rollback on any failed gate.",
    ]
    bullets = [p for p in document.paragraphs if p.style and p.style.name == "List Bullet"]
    for index, value in enumerate(rollout, 4):
        bullets[index].text = value

    for paragraph in document.paragraphs:
        if paragraph.text == "packaging/linux-client/AUDIT_PERSISTENCE_0_9_5.md":
            paragraph.text = "packaging/linux-client/AUDIT_PERSISTENCE_0_9_6.md"
    set_note(
        tables[11].cell(0, 0),
        "Decision",
        "The exact 0.9.6 archive is suitable only for the approved Ubuntu and Debian replacement compatibility pilots. Fleet expansion, production certification and broader platform acceptance remain blocked until the pending native gates pass.",
        True,
    )

    properties = document.core_properties
    properties.title = "Quiver seQuester Linux x64 0.9.6 Compatibility Candidate Addendum"
    properties.subject = "Immutable 0.9.6 explicit syscall-selector correction and pending replacement compatibility pilots"
    properties.author = "Quiver seQuester Documentation"
    properties.last_modified_by = "Quiver seQuester Documentation"
    properties.comments = ""

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Created {output}")


if __name__ == "__main__":
    main()
