from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from create_linux_092_release_addendum import GRAY, RED, add_field
from create_linux_095_release_addendum import set_note, set_row


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    document = Document(args.template)
    document.sections[0].header.paragraphs[0].text = "Quiver seQuester  |  Linux 0.9.6 compatibility acceptance"
    footer = document.sections[0].footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("RF DOC 001-G  •  13 September 2026  |  Page ")
    add_field(footer.add_run(), " PAGE ", "1")
    footer.add_run(" of ")
    add_field(footer.add_run(), " NUMPAGES ", "4")
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)

    replacements = {
        "Linux x64 0.9.6 Compatibility Candidate Addendum": "Linux x64 0.9.6 Compatibility Acceptance Addendum",
        "2. Explicit syscall selector correction": "2. Accepted compatibility results",
        "4. Verified pre-pilot evidence": "4. Verified native evidence",
        "5. Live acceptance gates": "5. Remaining acceptance gates",
        "6. Controlled pilot and rollback": "6. Controlled fleet expansion",
    }
    for paragraph in document.paragraphs:
        if paragraph.text in replacements:
            paragraph.text = replacements[paragraph.text]
        if "Replacement compatibility pilots pending" in paragraph.text:
            paragraph.text = "Published 13 September 2026  •  Two compatibility pilots accepted  •  Serial rollout in progress"
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(RED)

    tables = document.tables
    set_note(
        tables[0].cell(0, 0),
        "Supersession",
        "This addendum supersedes RF DOC 001-F only for current Linux x64 0.9.6 compatibility status. "
        "RF DOC 001-F remains the immutable publication-time pending-pilot record, and RF DOC 001-E remains immutable single-host 0.9.5 evidence.",
    )
    set_note(
        tables[1].cell(0, 0),
        "Acceptance boundary",
        "The rollout owner reports that the exact immutable 0.9.6 archive passed the approved Ubuntu 22/auditctl 3.0.7 and Debian 13.7/auditctl 4.0.2 replacement compatibility pilots. "
        "This is two-combination acceptance, not completed fleet deployment or production certification.",
        True,
    )
    set_row(tables[2], 1, ["Document identifier", "RF DOC 001-G"])
    set_row(tables[2], 2, ["Supersedes", "RF DOC 001-F current compatibility guidance only"])
    set_row(tables[2], 4, ["Scope", "Linux x64 0.9.6 Ubuntu and Debian compatibility acceptance"])
    set_row(tables[2], 5, ["Status", "Unsigned development candidate; two combinations accepted; serial rollout in progress"])

    package_rows = [
        ["Release", "ransomfree-client-0.9.6-dev.20260913T184639Z-08804ca0-linux-x64"],
        ["Archive", "ransomfree-client-0.9.6-dev.20260913T184639Z-08804ca0-linux-x64.tar.gz"],
        ["SHA-256", "501a4ec589b3af6f90c2eca6b4a7762b3f0c45c80961b21872310c567b77f02e"],
        ["Size and inventory", "85,053,726 bytes; 481 files verified"],
        ["Endpoint / relay", "Endpoint 0.9.6; unchanged Native AOT relay 0.9.1, protocol 1"],
        ["Approval", "Supported enrolled Linux fleet rollout in serial batches; not production certified"],
    ]
    for index, values in enumerate(package_rows, 1):
        set_row(tables[3], index, values)

    compatibility = [
        "Plex01 Ubuntu 22 and auditctl 3.0.7 upgraded from accepted 0.9.5: initial reconciliation was expected 6, already present 6, added 0.",
        "Nginx01 Debian 13.7 and auditctl 4.0.2 started from a clean 0.8.0 baseline without seQuester rules: initial reconciliation was expected 6, already present 0, added 6.",
        "Both hosts retained six unique explicit-all rules; actual auditd restarts caused one new loader invocation each and restored rules without manual rescue; intact loader reloads added 0.",
        "Both hosts passed root health, Full attribution, capability-free root collector and auditd-child relay, NoNewPrivileges, enrollment preservation and all-canary-hash checks.",
    ]
    bullets = [p for p in document.paragraphs if p.style and p.style.name == "List Bullet"]
    for index, value in enumerate(compatibility):
        bullets[index].text = value
    set_note(
        tables[4].cell(0, 0),
        "Ubuntu evidence",
        "Event 4a4417c238ecbdae98b1234912f7501c reached management sequence 1603 with /usr/bin/cat PID 211700, PPID 211643 and UID/eUID 0. "
        "The rollout owner reports exact Wazuh TCP receiver arrival from CPTNYC_DC01.",
    )
    set_note(
        tables[5].cell(0, 0),
        "Debian evidence",
        "Event 8604695684554e392a4e38b5da0df766 reached management and receiver checkpoint 1699 with /usr/bin/cat PID 2003918, PPID 2003868 and UID/eUID 0. "
        "The clean host's empty unrelated audit baseline remained unchanged.",
    )
    set_row(
        tables[6],
        1,
        [
            "auditd baseline",
            "Preserve each host's effective baseline: MDWE=yes on reviewed hosts; ProtectHome=yes on Ubuntu plex01 and no on Debian nginx01/backup01. No hardening override is authorized or required.",
        ],
    )

    evidence_rows = [
        ["Package integrity", "Passed", "Exact immutable archive and 481-file payload unchanged; SHA-256 verified."],
        ["Ubuntu reconciliation", "Passed", "Initial 6/6/0; six explicit-all rules; repeated real restart recovery; intact reload added 0."],
        ["Debian reconciliation", "Passed", "Clean initial 6/0/6; six explicit-all rules; repeated real restart recovery; intact reload added 0."],
        ["Attribution", "Passed", "Both controlled reads reached management with exact event, process, parent and UID/eUID attribution."],
        ["Security and preservation", "Passed", "Capability-free/NNP1 identities, per-host hardening, enrollment and all canary hashes preserved."],
    ]
    for index, values in enumerate(evidence_rows, 1):
        set_row(tables[7], index, values)
    set_note(
        tables[8].cell(0, 0),
        "Evidence ownership",
        "These live results were reported and recorded by the rollout owner. This documentation pass reconciled the approved evidence and did not independently operate either host or repeat the runtime tests.",
    )

    gate_rows = [
        ["Exact package", "Passed", "SHA-256 501a4ec589b3af6f90c2eca6b4a7762b3f0c45c80961b21872310c567b77f02e."],
        ["Ubuntu 22 / auditctl 3.0.7", "Passed", "Upgrade, idempotence, restart restoration, health, attribution and delivery reported accepted."],
        ["Debian 13.7 / auditctl 4.0.2", "Passed", "Clean integration, explicit-all adds, restart restoration, health, attribution and delivery reported accepted."],
        ["Supported fleet rollout", "In progress", "Rollout owner proceeding in serial batches; completion is not reported."],
        ["Unsupported / unreachable hosts", "Excluded", "tpot01 Debian 11 unsupported; itarian01 and misp01 unreachable."],
        ["Wazuh processing", "Pending", "Exact-event wire arrival passed; decoder, rule, index and dashboard acceptance remain open."],
        ["Broader native recovery", "Pending", "Reboot and nonempty unrelated live-policy preservation remain open."],
        ["Platform and UI", "Pending", "ARM64 relay, broader operation/filesystem/mount/immutable-policy matrix and UI acceptance."],
        ["Production certification", "Pending", "Two-combination compatibility acceptance does not certify production or complete the fleet."],
    ]
    for index, values in enumerate(gate_rows, 1):
        set_row(tables[9], index, values)

    set_note(
        tables[10].cell(0, 0),
        "Authoritative procedure",
        "Use packaging/linux-client/AUDIT_PERSISTENCE_0_9_6.md with the immutable baseline erratum, SHA-256 59cd4a04c8ba61eb280a62b5cac2bdf22836a18737c3d8ed87b6f0c42bc675cd. "
        "Preserve each host's effective ProtectHome state; do not alter the archive or extracted manifests.",
    )
    rollout = [
        "Continue only to supported enrolled Linux hosts under the rollout owner's existing serial-batch authority; do not imply fleet completion.",
        "Require exact package identity, six explicit-all rules, one loader invocation per auditd restart and zero-add intact reload on every accepted host.",
        "Preserve endpoint identity, queues, cursor, inbox, canaries, unrelated audit policy and each host's recorded security baseline.",
        "Stop and use host-specific rollback on any reconciliation, health, preservation, attribution or delivery failure; keep unsupported and unreachable hosts excluded.",
    ]
    bullets = [p for p in document.paragraphs if p.style and p.style.name == "List Bullet"]
    for index, value in enumerate(rollout, 4):
        bullets[index].text = value

    set_note(
        tables[11].cell(0, 0),
        "Decision",
        "The exact 0.9.6 archive has passed the approved Ubuntu 22/auditctl 3.0.7 and Debian 13.7/auditctl 4.0.2 compatibility pilots. "
        "The rollout owner may continue supported enrolled Linux deployment in serial batches under existing authority. Fleet completion and production certification remain unclaimed.",
        True,
    )

    properties = document.core_properties
    properties.title = "Quiver seQuester Linux x64 0.9.6 Compatibility Acceptance Addendum"
    properties.subject = "Two-distribution compatibility acceptance and controlled serial rollout status"
    properties.author = "Quiver seQuester Documentation"
    properties.last_modified_by = "Quiver seQuester Documentation"
    properties.comments = ""

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Created {output}")


if __name__ == "__main__":
    main()
